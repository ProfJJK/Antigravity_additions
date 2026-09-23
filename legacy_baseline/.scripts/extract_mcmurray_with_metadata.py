import os
import glob
import json
import time
import subprocess
import traceback
import sys
from pathlib import Path
from pptx import Presentation as PPTXPresentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
import logging

# Ensure cp1252 exceptions are bypassed
sys.stdout.reconfigure(encoding='utf-8')

logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")
logger = logging.getLogger("mcmurray_extractor")

STATE_FILE = Path(r"D:\__CoChem\__agentic\.scripts\mcmurray_extraction_state.json")
LOG_FILE = Path(r"D:\__CoChem\__agentic\.scripts\.logs\mcmurray_extraction.log")

os.makedirs(LOG_FILE.parent, exist_ok=True)

def load_state():
    if STATE_FILE.exists():
        with open(STATE_FILE, "r", encoding="utf-8") as f:
            return json.load(f)
    return {"completed_files": [], "status": "RUNNING"}

def save_state(state):
    with open(STATE_FILE, "w", encoding="utf-8") as f:
        json.dump(state, f, indent=4)

def extract_metadata_via_subprocess(image_path, slide_context):
    prompt = f"Analyze the following slide text context: '{slide_context}'. Generate a short descriptive metadata label for the accompanying graphic. Limit to 1 sentence."
    cmd = ["agy", "--agent", "cochem-coder", "-p", prompt]
    try:
        proc = subprocess.run(creationflags=0x08000000, cmd, capture_output=True, text=True, check=True, timeout=60)
        return proc.stdout.strip()
    except subprocess.TimeoutExpired:
        logger.error(f"Timeout extracting metadata for {image_path}")
        return "[TIMEOUT] Metadata generation failed."
    except subprocess.CalledProcessError as e:
        logger.error(f"Subprocess error: {e}")
        return "[ERROR] Metadata generation failed."

def process_file(pptx_path, base_materials_dir):
    filename = os.path.basename(pptx_path)
    base_name = os.path.splitext(filename)[0]
    out_dir = os.path.join(base_materials_dir, base_name)
    os.makedirs(out_dir, exist_ok=True)
    
    prs = PPTXPresentation(pptx_path)
    slides_data = []
    
    for i, slide in enumerate(prs.slides):
        slide_num = i + 1
        title = f"Slide {slide_num}"
        bullets = []
        image_path = None
        speaker_notes = None
        img_full_path = None
        
        if slide.shapes.title and slide.shapes.title.has_text_frame:
            title_text = slide.shapes.title.text.strip()
            if title_text:
                title = title_text
                
        img_idx = 1
        for shape in slide.shapes:
            if shape == slide.shapes.title:
                continue
                
            if shape.has_text_frame:
                text = shape.text.strip()
                if text:
                    for t in text.split('\n'):
                        clean_t = t.strip()
                        if clean_t:
                            bullets.append(clean_t)
            
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                if not image_path:
                    try:
                        image = shape.image
                        image_filename = f"slide_{slide_num}_img_{img_idx}.{image.ext}"
                        img_full_path = os.path.join(out_dir, image_filename)
                        with open(img_full_path, "wb") as f:
                            f.write(image.blob)
                        
                        image_path = os.path.join("materials_mcmurray", base_name, image_filename).replace("\\", "/")
                        img_idx += 1
                    except Exception as e:
                        logger.error(f"Error extracting image on slide {slide_num}: {e}")
                        
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes_text = slide.notes_slide.notes_text_frame.text.strip()
            if notes_text:
                speaker_notes = notes_text
                
        metadata = "No graphic"
        if image_path and img_full_path:
            context = " ".join(bullets)
            logger.info(f"Generating metadata for {image_path} via AGY...")
            metadata = extract_metadata_via_subprocess(img_full_path, context)
            
        s = {
            "slide_number": slide_num,
            "title": title,
            "bullets": bullets,
            "image_path": image_path,
            "image_metadata": metadata,
            "speaker_notes": speaker_notes
        }
        slides_data.append(s)
        
    json_path = os.path.join(out_dir, f"{base_name}_extracted.json")
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump({"slides": slides_data}, f, indent=4)
        
    logger.info(f"Extracted {filename} to {out_dir}")

def main():
    input_dir = sys.argv[1] if len(sys.argv) > 1 else r"D:\Gdrive\_CU Teaching\2026FA CHEM103 Fundamentals of Chemistry\GOB_McMurray"
    materials_dir = r"D:\Gdrive\_CU Teaching\2026FA CHEM103 Fundamentals of Chemistry\materials_mcmurray"
    
    os.makedirs(materials_dir, exist_ok=True)
    pptx_files = sorted(glob.glob(os.path.join(input_dir, "*.pptx")))
    
    state = load_state()
    if state["status"] == "COMPLETED":
        logger.info("All files already completed according to state.")
        return

    for pptx_file in pptx_files:
        filename = os.path.basename(pptx_file)
        if filename in state["completed_files"]:
            continue
            
        logger.info(f"Processing {filename}...")
        try:
            process_file(pptx_file, materials_dir)
            state["completed_files"].append(filename)
            save_state(state)
        except Exception as e:
            logger.error(f"Critical error on {filename}: {e}\n{traceback.format_exc()}")
            state["status"] = "ERROR"
            save_state(state)
            break
            
    if len(state["completed_files"]) == len(pptx_files):
        state["status"] = "COMPLETED"
        save_state(state)
        logger.info("Extraction State Machine Finished Successfully.")

if __name__ == "__main__":
    main()
