import collections 
import collections.abc
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
import os
import sys

def extract_pptx(filepath):
    target_dir = os.path.dirname(filepath)
    base_name = os.path.splitext(os.path.basename(filepath))[0]
    
    assets_dir = os.path.join(target_dir, f"{base_name}_assets")
    os.makedirs(assets_dir, exist_ok=True)
    
    outline_path = os.path.join(target_dir, f"{base_name}_outline.md")
    
    print(f"[*] Extracting legacy presentation: {filepath}")
    
    try:
        prs = Presentation(filepath)
    except Exception as e:
        print(f"[!] Error opening PPTX: {e}")
        print("Ensure 'python-pptx' is installed: pip install python-pptx")
        sys.exit(1)
        
    image_count = 0
    
    with open(outline_path, 'w', encoding='utf-8') as f:
        f.write(f"# Extracted Outline: {base_name}\n\n")
        
        for i, slide in enumerate(prs.slides):
            f.write(f"## Slide {i+1}\n\n")
            
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    f.write(shape.text.strip() + "\n\n")
                    
                # Extract Images
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    try:
                        image = shape.image
                        image_bytes = image.blob
                        image_ext = image.ext
                        image_count += 1
                        image_filename = f"slide_{i+1}_img_{image_count}.{image_ext}"
                        image_path = os.path.join(assets_dir, image_filename)
                        
                        with open(image_path, "wb") as img_f:
                            img_f.write(image_bytes)
                            
                        f.write(f"![Extracted Image]({base_name}_assets/{image_filename})\n\n")
                    except Exception as e:
                        print(f"[-] Could not extract image on slide {i+1}: {e}")
                    
            # Extract Speaker Notes
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    f.write(f"**Speaker Notes:**\n{notes}\n\n")
                
            f.write("---\n\n")
            
    print(f"[*] Extraction complete. Found {image_count} images. Outline saved to {outline_path}")
    return outline_path, assets_dir

if __name__ == "__main__":
    if len(sys.argv) > 1:
        extract_pptx(sys.argv[1])
    else:
        print("Usage: python pptx_extractor.py <path_to_pptx>")
